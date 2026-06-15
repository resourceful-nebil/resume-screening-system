from io import BytesIO
from pathlib import Path

import PyPDF2
from docx import Document

from extractors.keyword_extractor import KeywordExtractor


class ResumeParser:
    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt"}

    def __init__(self):
        self.extractor = KeywordExtractor()

    def parse_file(self, file_path: str | Path) -> dict:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Resume file not found: {path}")

        suffix = path.suffix.lower()
        if suffix == ".pdf":
            text = self._extract_pdf(path)
        elif suffix in {".docx", ".doc"}:
            text = self._extract_docx(path)
        elif suffix == ".txt":
            text = path.read_text(encoding="utf-8", errors="ignore").strip()
        else:
            raise ValueError(f"Unsupported file type: {suffix}")

        contact = self.extractor.extract_contact_info(text)
        return {
            "file_name": path.name,
            "file_path": str(path),
            "text": text,
            "name": contact["name"],
            "email": contact["email"],
            "phone": contact["phone"],
        }

    def parse(self, source: str | Path | bytes, filename: str = "") -> str:
        if isinstance(source, bytes):
            return self._parse_bytes(source, filename)
        return self.parse_file(source)["text"]

    def parse_upload(self, data: bytes, filename: str) -> dict:
        text = self._parse_bytes(data, filename)
        contact = self.extractor.extract_contact_info(text)
        return {
            "file_name": filename,
            "text": text,
            "name": contact["name"],
            "email": contact["email"],
            "phone": contact["phone"],
            "file_bytes": data,
        }

    def _parse_bytes(self, data: bytes, filename: str) -> str:
        suffix = Path(filename).suffix.lower()
        if suffix == ".pdf":
            reader = PyPDF2.PdfReader(BytesIO(data))
            text = ""
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
            return text.strip()
        if suffix in {".docx", ".doc"}:
            doc = Document(BytesIO(data))
            return "\n".join(para.text for para in doc.paragraphs).strip()
        if suffix == ".txt":
            return data.decode("utf-8", errors="ignore").strip()
        raise ValueError(f"Unsupported file type: {suffix}")

    def _extract_pdf(self, file_path: Path) -> str:
        text = ""
        with open(file_path, "rb") as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        return text.strip()

    def _extract_docx(self, file_path: Path) -> str:
        doc = Document(file_path)
        return "\n".join(para.text for para in doc.paragraphs).strip()
